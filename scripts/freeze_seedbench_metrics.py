"""Freeze reviewed SeedBench summaries without rerunning any evaluator."""

from __future__ import annotations

import argparse
from contextlib import ExitStack, closing
import hashlib
import json
import math
from numbers import Real
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd


OUTPUT_FILENAMES = {
    "metrics": "SeedBench-for_plot.tsv",
    "provenance": "SeedBench-for_plot-provenance.json",
}
MODEL_ORDER = (
    "Phyto-Reasoner", "GPT-5", "o3", "Gemini-2.5-Pro",
    "Claude-Opus-4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
)
SUMMARY_HEADERS = (
    "Phyto-Reasoner", "GPT-5", "o3", "Gemini 2.5 Pro",
    "Claude Opus 4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
)
DETAIL_HEADERS = (
    "Phyto_Reasoner", "GPT-5", "o3", "Gemini-2.5-pro",
    "Claude-Opus-4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
)
EXPECTED_COUNTS = {"QA-1": 235, "QA-2": 218, "QA-3": 269, "QA-4": 277}
METRICS = {
    "QA-1": "Accuracy", "QA-2": "Macro-F1",
    "QA-3": "ROUGE-L F1", "QA-4": "ROUGE-L F1",
}
ROUGE_PATTERN = re.compile(
    r"ROUGE-L F1 score=([+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)"
)
SUMMARY_DECIMALS = 4
SUMMARY_ABS_TOLERANCE = 1e-12


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_workbook(path: Path) -> dict[str, list[tuple]]:
    """Read each sheet once per mode; retain no prompts or responses."""
    sheets = {}
    with ExitStack() as stack:
        cached = stack.enter_context(closing(load_workbook(path, read_only=True, data_only=True)))
        formulas = stack.enter_context(closing(load_workbook(path, read_only=True, data_only=False)))
        for name in ("summary", *EXPECTED_COUNTS):
            if name not in cached.sheetnames:
                raise ValueError(f"Missing required sheet {name}.")
            width = 10 if name == "summary" else 25
            rows = []
            value_rows = cached[name].iter_rows(max_col=width)
            formula_rows = formulas[name].iter_rows(max_col=width)
            for row_number, (values, expressions) in enumerate(zip(value_rows, formula_rows, strict=True), 1):
                indices = range(width) if name == "summary" or row_number <= 2 else (0, *range(10, 25, 2))
                selected = []
                for index in indices:
                    value = values[index].value
                    if expressions[index].data_type == "f" and value is None:
                        coordinate = f"{get_column_letter(index + 1)}{row_number}"
                        raise ValueError(f"{name}!{coordinate} has no cached formula value.")
                    selected.append(value)
                rows.append(tuple(selected))
            sheets[name] = rows
    return sheets


def _score(value: object, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, Real):
        raise ValueError(f"{label} must contain a finite numeric score in [0, 1].")
    result = float(value)
    if not math.isfinite(result) or not 0 <= result <= 1:
        raise ValueError(f"{label} must contain a finite numeric score in [0, 1].")
    return result


def _validate_detail(task: str, rows: list[tuple]) -> list[tuple]:
    expected_count = EXPECTED_COUNTS[task]
    if len(rows) != expected_count + 2:
        raise ValueError(f"{task} requires exactly {expected_count} item rows; found {len(rows) - 2}.")
    if rows[0][0] != "ID":
        raise ValueError(f"{task}!A1 must be ID.")
    for index, header in enumerate(DETAIL_HEADERS):
        column = 9 + 2 * index
        if rows[0][column] != header:
            raise ValueError(f"{task}!{get_column_letter(column + 1)}1 must be {header!r}.")
        if rows[1][column:column + 2] != ("Response", "Final judgment result"):
            raise ValueError(f"{task} has invalid response/judgment headers for {MODEL_ORDER[index]}.")
    seen = set()
    for row_number, row in enumerate(rows[2:], 3):
        question_id = row[0]
        if not isinstance(question_id, str) or not question_id.strip():
            raise ValueError(f"{task}!A{row_number} has a missing or invalid ID.")
        if question_id in seen:
            raise ValueError(f"{task}!A{row_number} has a duplicate ID.")
        seen.add(question_id)
    return rows[2:]


def freeze_metrics(source_path: Path) -> tuple[pd.DataFrame, dict[str, object]]:
    source_path = Path(source_path)
    sheets = _read_workbook(source_path)
    summary = sheets["summary"]
    if len(summary) != 5 or summary[0] != ("task_id", "metric", *SUMMARY_HEADERS):
        raise ValueError("summary must have the current task/metric/model headers and four task rows.")

    records = []
    rouge_aggregation = {}
    for row_index, (task, count) in enumerate(EXPECTED_COUNTS.items(), 1):
        summary_row = summary[row_index]
        if summary_row[:2] != (task, METRICS[task]):
            raise ValueError(f"summary row {row_index + 1} must be {task} with metric {METRICS[task]}.")
        detail = _validate_detail(task, sheets[task])
        if task in ("QA-3", "QA-4"):
            rouge_aggregation[task] = {}
        for model_index, model in enumerate(MODEL_ORDER):
            summary_cell = f"summary!{get_column_letter(model_index + 3)}{row_index + 1}"
            score = _score(summary_row[model_index + 2], summary_cell)
            if task in ("QA-3", "QA-4"):
                scores = []
                column = get_column_letter(11 + 2 * model_index)
                for row_number, row in enumerate(detail, 3):
                    judgment = row[model_index + 1]
                    label = f"{task}!{column}{row_number}"
                    match = ROUGE_PATTERN.fullmatch(judgment.strip()) if isinstance(judgment, str) else None
                    if match is None:
                        raise ValueError(f"{label} must contain exactly one ROUGE-L F1 score log.")
                    scores.append(_score(float(match.group(1)), label))
                mean = math.fsum(scores) / count
                rounded_mean = round(mean, SUMMARY_DECIMALS)
                if not math.isclose(score, rounded_mean, rel_tol=0, abs_tol=SUMMARY_ABS_TOLERANCE):
                    raise ValueError(f"{summary_cell} disagrees with the four-decimal all-item mean for {task}/{model}.")
                score = rounded_mean
                rouge_aggregation[task][model] = {
                    "mean_unrounded": mean,
                    "zero_count": scores.count(0.0),
                }
            records.append({"Task": task, "Metric": METRICS[task], "Model": model, "Score": score, "N": count})

    provenance = {
        "schema_version": 1,
        "source_sha256": _sha256(source_path),
        "source_rows": dict(EXPECTED_COUNTS),
        "model_order": list(MODEL_ORDER),
        "task_order": list(EXPECTED_COUNTS),
        "source_columns": {
            model: {
                "summary_header": SUMMARY_HEADERS[index],
                "summary_column": get_column_letter(index + 3),
                "detail_header": DETAIL_HEADERS[index],
                "judgment_column": get_column_letter(11 + 2 * index),
            }
            for index, model in enumerate(MODEL_ORDER)
        },
        "metric_definitions": {
            "QA-1": "Accuracy; supplied summary values retained unchanged without reevaluation.",
            "QA-2": (
                "Macro-F1; harmonic mean of sample-averaged precision and sample-averaged recall, "
                "not the mean of per-item F1. Supplied summary values retained unchanged without reevaluation."
            ),
            "QA-3": "Arithmetic mean of all existing per-item ROUGE-L F1 logs, including zeros; equal weights, no filtering or reevaluation.",
            "QA-4": "Arithmetic mean of all existing per-item ROUGE-L F1 logs, including zeros; equal weights, no filtering or reevaluation.",
        },
        "summary_reconciliation": {
            "tasks": ["QA-3", "QA-4"],
            "decimal_places": SUMMARY_DECIMALS,
            "rounding": "Round the complete arithmetic mean to four decimals, preserving the supplied summary/plot precision.",
            "absolute_tolerance": SUMMARY_ABS_TOLERANCE,
        },
        "rouge_aggregation": rouge_aggregation,
    }
    return pd.DataFrame(records, columns=("Task", "Metric", "Model", "Score", "N")), provenance


def write_outputs(table: pd.DataFrame, provenance: dict[str, object], output_dir: Path) -> None:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / OUTPUT_FILENAMES["metrics"]
    table.to_csv(path, sep="\t", index=False, float_format="%.15g", lineterminator="\n")
    finalized = {**provenance, "outputs": {path.name: _sha256(path)}}
    (output_dir / OUTPUT_FILENAMES["provenance"]).write_text(
        json.dumps(finalized, indent=2, sort_keys=True) + "\n", encoding="utf-8",
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    table, provenance = freeze_metrics(args.source)
    write_outputs(table, provenance, args.output_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
