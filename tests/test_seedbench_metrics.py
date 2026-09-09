import hashlib
import importlib
import json
import math
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook


MODEL_ORDER = [
    "Phyto-Reasoner", "GPT-5", "o3", "Gemini-2.5-Pro",
    "Claude-Opus-4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
]
SUMMARY_HEADERS = [
    "Phyto-Reasoner", "GPT-5", "o3", "Gemini 2.5 Pro",
    "Claude Opus 4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
]
DETAIL_HEADERS = [
    "Phyto_Reasoner", "GPT-5", "o3", "Gemini-2.5-pro",
    "Claude-Opus-4.1", "DeepSeek-R1", "DeepSeek-V3", "Grok-3-Beta",
]
COUNTS = {"QA-1": 235, "QA-2": 218, "QA-3": 269, "QA-4": 277}
METRICS = {"QA-1": "Accuracy", "QA-2": "Macro-F1",
           "QA-3": "ROUGE-L F1", "QA-4": "ROUGE-L F1"}


def _scores(task, model_index):
    return [0.0] + [(model_index + 1) / 10] * (COUNTS[task] - 1)


@pytest.fixture
def workbook_path(tmp_path):
    path = tmp_path / "private-source-sentinel.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["task_id", "metric", *SUMMARY_HEADERS])
    for task in COUNTS:
        if task == "QA-1":
            scores = [0.1234 + i / 100 for i in range(8)]
        elif task == "QA-2":
            scores = [0.4321 + i / 100 for i in range(8)]
        else:
            scores = [round(math.fsum(_scores(task, i)) / COUNTS[task], 4)
                      for i in range(8)]
        summary.append([task, METRICS[task], *scores])
        sheet = workbook.create_sheet(task)
        sheet.cell(1, 1, "ID")
        for i, model in enumerate(DETAIL_HEADERS):
            column = 10 + 2 * i
            sheet.cell(1, column, model)
            sheet.cell(2, column, "Response")
            sheet.cell(2, column + 1, "Final judgment result")
        for i in range(COUNTS[task]):
            row = i + 3
            sheet.cell(row, 1, f"private-id-{task}-{i}")
            sheet.cell(row, 4, "private-question-sentinel")
            for model_index in range(8):
                column = 10 + 2 * model_index
                sheet.cell(row, column, "private-response-sentinel")
                score = _scores(task, model_index)[i]
                judgment = (f"ROUGE-L F1 score={score:.6f}" if task in ("QA-3", "QA-4")
                            else "preserve-summary-not-per-item-f1")
                sheet.cell(row, column + 1, judgment)
    workbook.save(path)
    workbook.close()
    return path


def _edit(path, sheet, coordinate, value):
    workbook = load_workbook(path)
    workbook[sheet][coordinate] = value
    workbook.save(path)
    workbook.close()


def _freezer():
    return importlib.import_module("scripts.freeze_seedbench_metrics")


def test_freeze_preserves_summary_scores_and_averages_all_rouge_logs(workbook_path):
    table, provenance = _freezer().freeze_metrics(workbook_path)
    assert list(table.columns) == ["Task", "Metric", "Model", "Score", "N"]
    assert list(zip(table.Task, table.Model)) == [
        (task, model) for task in COUNTS for model in MODEL_ORDER
    ]
    for task, group in table.groupby("Task", sort=False):
        assert group.N.tolist() == [COUNTS[task]] * 8
        assert group.Metric.tolist() == [METRICS[task]] * 8
        if task in ("QA-1", "QA-2"):
            base = 0.1234 if task == "QA-1" else 0.4321
            assert group.Score.tolist() == pytest.approx([base + i / 100 for i in range(8)])
        else:
            for i, score in enumerate(group.Score):
                mean = math.fsum(_scores(task, i)) / COUNTS[task]
                assert score == round(mean, 4)
                details = provenance["rouge_aggregation"][task][MODEL_ORDER[i]]
                assert details["mean_unrounded"] == pytest.approx(mean)
                assert details["zero_count"] == 1
    assert "sample-averaged precision" in provenance["metric_definitions"]["QA-2"]
    assert "not the mean of per-item F1" in provenance["metric_definitions"]["QA-2"]
    assert provenance["source_sha256"] == hashlib.sha256(workbook_path.read_bytes()).hexdigest()


@pytest.mark.parametrize("judgment", [
    None, "", "ROUGE-L F1 score=NaN", "ROUGE-L F1 score=inf",
    "ROUGE-L F1 score=-0.1", "ROUGE-L F1 score=1.1",
    "ROUGE-L precision score=0.5", "ROUGE-L F1 score=0.5; score=0.9",
    "ROUGE-L F1 score=0.5 garbage", 0.5,
])
def test_freeze_rejects_missing_invalid_or_ambiguous_scores(workbook_path, judgment):
    _edit(workbook_path, "QA-3", "K3", judgment)
    with pytest.raises(ValueError, match="QA-3.*K3"):
        _freezer().freeze_metrics(workbook_path)


@pytest.mark.parametrize("task", COUNTS)
@pytest.mark.parametrize("value", [None, "private-id-placeholder"])
def test_freeze_rejects_missing_and_duplicate_ids(workbook_path, task, value):
    if value is not None:
        value = f"private-id-{task}-0"
    _edit(workbook_path, task, "A4", value)
    with pytest.raises(ValueError, match=f"{task}.*ID"):
        _freezer().freeze_metrics(workbook_path)


@pytest.mark.parametrize("task", COUNTS)
@pytest.mark.parametrize("delta", [-1, 1])
def test_freeze_requires_the_full_fixed_cohort(workbook_path, task, delta):
    workbook = load_workbook(workbook_path)
    sheet = workbook[task]
    if delta == -1:
        sheet.delete_rows(COUNTS[task] + 2)
    else:
        sheet.cell(COUNTS[task] + 3, 1, "unexpected-extra-private-id")
    workbook.save(workbook_path)
    workbook.close()
    with pytest.raises(ValueError, match=f"{task}.*{COUNTS[task]}"):
        _freezer().freeze_metrics(workbook_path)


@pytest.mark.parametrize(("sheet", "coordinate", "value"), [
    ("summary", "C4", 0.4), ("summary", "C5", 0.4),
    ("summary", "G1", "Claude-4.1-Opus"),
    ("QA-3", "R1", "Claude-4.1-Opus"),
    ("QA-4", "P1", "Gemini 2.5 Pro"),
    ("QA-4", "K2", "Score"),
    ("summary", "A4", "QA-4"), ("summary", "B4", "ROUGE-L precision"),
])
def test_freeze_rejects_stale_summary_or_unrecognized_headers(workbook_path, sheet, coordinate, value):
    _edit(workbook_path, sheet, coordinate, value)
    with pytest.raises(ValueError, match=sheet):
        _freezer().freeze_metrics(workbook_path)


def test_freeze_rejects_a_formula_without_a_cached_value(workbook_path):
    _edit(workbook_path, "summary", "C4", "=AVERAGE(0,0.1)")
    with pytest.raises(ValueError, match="summary.*C4.*cached"):
        _freezer().freeze_metrics(workbook_path)


def test_written_outputs_are_deterministic_and_do_not_expose_private_content(workbook_path, tmp_path):
    module = _freezer()
    table, provenance = module.freeze_metrics(workbook_path)
    assert "outputs" not in provenance
    output_dir = tmp_path / "frozen"
    module.write_outputs(table, provenance, output_dir)
    tsv = output_dir / "SeedBench-for_plot.tsv"
    metadata = output_dir / "SeedBench-for_plot-provenance.json"
    original_bytes = (tsv.read_bytes(), metadata.read_bytes())
    module.write_outputs(table, provenance, output_dir)
    assert original_bytes == (tsv.read_bytes(), metadata.read_bytes())
    assert "outputs" not in provenance
    loaded = json.loads(metadata.read_text())
    assert loaded["outputs"] == {tsv.name: hashlib.sha256(tsv.read_bytes()).hexdigest()}
    pd.testing.assert_frame_equal(table, pd.read_csv(tsv, sep="\t"))
    for forbidden in (str(workbook_path), workbook_path.name, "private-id-",
                      "private-question-sentinel", "private-response-sentinel"):
        assert forbidden not in metadata.read_text()
        assert forbidden not in tsv.read_text()


def test_cli_uses_the_explicit_source_and_output_directory(workbook_path, tmp_path):
    output_dir = tmp_path / "cli-output"
    assert _freezer().main(["--source", str(workbook_path), "--output-dir", str(output_dir)]) == 0
    assert sorted(path.name for path in output_dir.iterdir()) == [
        "SeedBench-for_plot-provenance.json", "SeedBench-for_plot.tsv",
    ]
